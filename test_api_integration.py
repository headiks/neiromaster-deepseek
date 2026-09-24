"""
Интеграционные тесты API на настоящем приложении и настоящей PostgreSQL.

Поднимают FastAPI-приложение целиком (lifespan: схема, миграции, владелец) через
TestClient и проходят сценарии глазами пользователей разных ролей: вход, права
по подразделениям, штатка, планы, документы, вопросы, заголовки безопасности.
DeepSeek не нужен: сценарии не зовут модель (или проверяют, что до неё не доходит).

Нужна ОТДЕЛЬНАЯ тестовая БД — схема public в ней пересоздаётся:
    NEIROMASTER_TEST_DSN (по умолчанию postgresql://neiromaster:neiromaster@localhost:5432/neiromaster_test)
Нет БД — тесты пропускаются. Запуск: python -m pytest -q test_api_integration.py
"""
import io
import os
import sys
import shutil
import tempfile
from pathlib import Path

import pytest

TEST_DSN = (os.environ.get("NEIROMASTER_TEST_DSN")
            or "postgresql://neiromaster:neiromaster@localhost:5432/neiromaster_test")

# Окружение — до импорта модулей приложения (они читают его при импорте).
os.environ.update({
    "NEIROMASTER_DB_DSN": TEST_DSN,
    "NEIROMASTER_INSECURE_COOKIE": "1",       # TestClient ходит по http
    "NEIROMASTER_SCHEDULER": "0",
    "REDIS_URL": "",                          # всё в памяти процесса — тест самодостаточен
    "DEEPSEEK_API_KEY": "",
    "NEIROMASTER_ADMIN_PASSWORD": "owner-initial-pass",
    "NEIROMASTER_PII_KEY": "",
})

try:
    import psycopg
    with psycopg.connect(TEST_DSN, connect_timeout=3) as _c:
        _c.execute("DROP SCHEMA public CASCADE")
        _c.execute("CREATE SCHEMA public")
    DB_OK = True
except Exception as _e:              # noqa: BLE001 — любая причина = нет БД
    DB_OK = False
    DB_ERR = str(_e)

pytestmark = pytest.mark.skipif(not DB_OK, reason="тестовая PostgreSQL недоступна")

BASE = Path(__file__).resolve().parent
ORIGIN = {"Origin": "http://testserver"}


@pytest.fixture(scope="module")
def env():
    tmp = Path(tempfile.mkdtemp())
    import users
    import indexing
    import config
    users.DATA_DIR = tmp
    users.INITIAL_CREDENTIALS_PATH = tmp / "owner_initial_credentials.txt"
    users.USERS_PATH = tmp / "users.json"
    users.LEGACY_EMPLOYEES_PATH = tmp / "employees.json"
    indexing.DOCS_DIR = tmp / "documents"
    indexing.DOCS_DIR.mkdir()
    config.REGISTRY_PATH = tmp / "registry.json"
    import docregistry, questions
    docregistry.REGISTRY_PATH = tmp / "registry.json"
    questions.QUESTIONS_PATH = tmp / "pending_questions.json"
    # фоновая индексация без docling/DeepSeek: только отметка, что задача поставлена
    import jobs
    jobs.enqueue_index = lambda job_id: True
    from fastapi.testclient import TestClient
    import app
    with TestClient(app.app) as client:
        yield {"client": client, "tmp": tmp}
    shutil.rmtree(tmp, ignore_errors=True)


def _login(client, username, password):
    client.cookies.clear()
    r = client.post("/api/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def owner(env):
    """Владелец: вход по начальному паролю -> свой пароль -> новая сессия."""
    import users
    c = env["client"]
    own = users.get_owner()
    first = _login(c, own["username"], "owner-initial-pass")
    assert first["must_change_credentials"]
    assert c.get("/admin", follow_redirects=False).headers["location"] == "/setup"
    r = c.post("/api/setup-credentials", json={"password": "owner-strong-pass"}, headers=ORIGIN)
    assert r.status_code == 200, r.text
    _login(c, own["username"], "owner-strong-pass")
    return {"username": own["username"], "password": "owner-strong-pass", "id": own["id"]}


def _as_owner(env, owner):
    _login(env["client"], owner["username"], owner["password"])
    return env["client"]


# ---------------------------------------------------------------- безопасность
def test_security_headers_and_healthz(env):
    r = env["client"].get("/healthz")
    assert r.status_code == 200 and r.json()["db"] is True
    h = r.headers
    assert h["x-frame-options"] == "DENY" and h["x-content-type-options"] == "nosniff"
    assert "frame-ancestors 'none'" in h["content-security-policy"]
    assert "unpkg" not in h["content-security-policy"]


def test_static_pages_have_no_external_scripts(env):
    for name in ("admin.html", "index.html", "login.html", "setup.html", "register.html"):
        text = (BASE / "static" / name).read_text(encoding="utf-8")
        assert "unpkg.com" not in text and "googleapis" not in text, name


def test_cross_site_post_with_cookie_rejected(env, owner):
    c = _as_owner(env, owner)
    r = c.post("/plans/template", headers={"Origin": "https://evil.example"})
    assert r.status_code == 403
    r = c.post("/plans/template", headers=ORIGIN)
    assert r.status_code == 200


def test_session_token_stored_hashed_and_legacy_upgraded(env, owner):
    import db
    import auth
    token, user = auth.login(owner["username"], owner["password"])
    assert db.query("SELECT 1 FROM sessions WHERE token = %s", (token,), "one") is None
    assert auth.get_session_user(token)["id"] == user["id"]
    # сессия «старого образца» (открытый токен) работает и сразу переводится на хэш
    import time
    db.execute("INSERT INTO sessions (token, user_id, created_at, seen_at) VALUES (%s,%s,%s,%s)",
               ("legacy-plain-token", user["id"], time.time(), time.time()))
    assert auth.get_session_user("legacy-plain-token")["id"] == user["id"]
    assert db.query("SELECT 1 FROM sessions WHERE token = 'legacy-plain-token'", (), "one") is None


def test_bearer_logout_invalidates_token(env, owner):
    c = env["client"]
    c.cookies.clear()
    tok = c.post("/api/login", json={"username": owner["username"], "password": owner["password"]}).json()["token"]
    c.cookies.clear()
    hdr = {"Authorization": f"Bearer {tok}"}
    assert c.get("/api/me", headers=hdr).status_code == 200
    assert c.post("/api/logout", headers=hdr).status_code == 200
    assert c.get("/api/me", headers=hdr).status_code == 401


def test_login_lockout_and_admin_reset(env, owner):
    import auth
    c = _as_owner(env, owner)
    emp = c.post("/users", json={"full_name": "Блокируемый Борис Борисович"}, headers=ORIGIN).json()
    for _ in range(auth.LOCKOUT_ATTEMPTS):
        c.post("/api/login", json={"username": emp["username"], "password": "wrong-password"})
    r = c.post("/api/login", json={"username": emp["username"], "password": emp["temp_password"]})
    assert r.status_code == 401 and "много" in r.json()["detail"].lower()
    c = _as_owner(env, owner)
    new = c.post(f"/users/{emp['id']}/credentials", json={}, headers=ORIGIN).json()["temp_password"]
    assert c.post("/api/login", json={"username": emp["username"], "password": new}).status_code == 200


# ---------------------------------------------------------------- люди и права
def _make_admin(env, owner, name, dept):
    c = _as_owner(env, owner)
    u = c.post("/users", json={"full_name": name, "department": dept}, headers=ORIGIN).json()
    assert c.post(f"/users/{u['id']}/role", json={"role": "admin"}, headers=ORIGIN).status_code == 200
    pw = c.post(f"/users/{u['id']}/credentials", json={"password": "admin-pass-123"}, headers=ORIGIN).json()
    c.cookies.clear()
    c.post("/api/login", json={"username": pw["username"], "password": "admin-pass-123"})
    assert c.post("/api/setup-credentials", json={"password": "admin-pass-456"}, headers=ORIGIN).status_code == 200
    return {"id": u["id"], "username": pw["username"], "password": "admin-pass-456"}


def test_department_admin_boundaries(env, owner):
    c = _as_owner(env, owner)
    other = c.post("/users", json={"full_name": "Чужой Иван Петрович", "department": "Склад"}, headers=ORIGIN).json()
    adm = _make_admin(env, owner, "Админов Олег Олегович", "Цех 1")
    _login(c, adm["username"], adm["password"])
    # создаёт только в своё подразделение
    mine = c.post("/users", json={"full_name": "Свой Пётр Петрович", "department": "Склад"}, headers=ORIGIN).json()
    assert mine["department"] == "Цех 1"
    # чужого не видит, не правит и не удаляет
    ids = {u["id"] for u in c.get("/users").json()["users"]}
    assert mine["id"] in ids and other["id"] not in ids
    assert c.delete(f"/users/{other['id']}", headers=ORIGIN).status_code == 403
    r = c.post("/users/bulk-delete", json={"ids": [other["id"]]}, headers=ORIGIN).json()
    assert r["deleted"] == 0 and r["skipped"]
    # не может «переехать» в другой отдел, чтобы увидеть его людей
    me = c.get(f"/users/{adm['id']}").json()
    c.put(f"/users/{adm['id']}", json={**me, "department": "Склад"}, headers=ORIGIN)
    assert c.get(f"/users/{adm['id']}").json()["department"] == "Цех 1"
    # своего — может удалить
    assert c.delete(f"/users/{mine['id']}", headers=ORIGIN).status_code == 200


def test_staffing_import_dates_and_foreign_departments(env, owner):
    adm = _make_admin(env, owner, "Кадров Кирилл Кириллович", "Цех 2")
    c = env["client"]
    _login(c, adm["username"], adm["password"])
    res = c.post("/staffing/import", json={"records": [
        {"full_name": "Датов Дмитрий Дмитриевич", "position": "Слесарь", "department": "", "start_date": "15.05.2026"},
        {"full_name": "Лишний Лев Львович", "position": "Токарь", "department": "Склад", "start_date": ""},
    ]}, headers=ORIGIN).json()
    assert [p["full_name"] for p in res["profiles"]] == ["Датов Дмитрий Дмитриевич"]
    assert any("подразделение" in s["reason"] for s in res["skipped"])
    import users
    u = users.get_by_username(res["profiles"][0]["username"])
    assert u["start_date"] == "2026-05-15" and u["department"] == "Цех 2"


def test_credentials_xlsx_formula_injection(env, owner):
    c = _as_owner(env, owner)
    u = c.post("/users", json={"full_name": "=HYPERLINK(\"http://evil\",\"x\")"}, headers=ORIGIN).json()
    body = c.get(f"/users/credentials.xlsx?ids={u['id']}").content
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(body)).active
    cell = ws.cell(row=2, column=1)
    assert cell.data_type == "s" and str(cell.value).startswith("'=")


# ---------------------------------------------------------------- планы
def test_plan_ids_sanitized_and_locking(env, owner):
    c = _as_owner(env, owner)
    evil = "x');alert(1);//"
    plan = c.post("/plans", headers=ORIGIN, json={"title": "Тест", "stages": [
        {"id": evil, "title": "Этап", "duration": {"value": 10 ** 9, "unit": "days"},
         "substages": [{"id": evil, "title": "Подэтап", "brief": "б"}]}]}).json()
    st = plan["stages"][0]
    assert "'" not in st["id"] and ")" not in st["substages"][0]["id"]
    assert st["duration"]["value"] <= 730
    # оптимистичная блокировка: сохранение поверх чужой правки -> 409
    pid, stamp = plan["plan_id"], plan["updated_at"]
    import time
    time.sleep(1.1)
    ok = c.put(f"/plans/{pid}", headers=ORIGIN, json={**plan, "title": "Правка 1", "expected_updated_at": stamp})
    assert ok.status_code == 200
    stale = c.put(f"/plans/{pid}", headers=ORIGIN, json={**plan, "title": "Правка 2", "expected_updated_at": stamp})
    assert stale.status_code == 409
    forced = c.put(f"/plans/{pid}", headers=ORIGIN, json={**plan, "title": "Правка 2"})
    assert forced.status_code == 200 and forced.json()["title"] == "Правка 2"


def test_standard_template_spreads_messages(env, owner):
    c = _as_owner(env, owner)
    plan = c.post("/plans/template", headers=ORIGIN).json()
    pre = plan["stages"][0]
    days = [s["schedule"]["day"] for s in pre["substages"]]
    assert max(days) > 1, "подэтапы стандартного плана разнесены по дням этапа"
    import planner
    per_day = {}
    for item in planner.resolve_schedule(plan):
        per_day[item["schedule"]["offset_days"]] = per_day.get(item["schedule"]["offset_days"], 0) + 1
    assert per_day[-3] < len(pre["substages"])


def test_group_daily_session(env, owner):
    import planner
    plan = planner.normalize_plan({"title": "t", "group_daily": True, "start_date": "2026-10-01", "stages": [
        {"title": "Первый день", "duration": {"value": 1, "unit": "days"}, "substages": [
            {"title": "a", "schedule": {"day": 1, "time": "10:00"}},
            {"title": "b", "schedule": {"day": 1, "time": "15:00"}}]}]})
    times = {i["schedule"]["send_at"] for i in planner.resolve_schedule(plan)}
    assert times == {"2026-10-01T10:00"}


# ---------------------------------------------------------------- документы
def test_upload_confidential_dedup_and_traversal(env, owner):
    c = _as_owner(env, owner)
    body = "Положение об оплате труда. Оклад 100 000 рублей.".encode("utf-8")
    r = c.post("/documents/upload?confidential=true", headers=ORIGIN,
               files={"file": ("Положение об оплате (2).txt", body, "text/plain")})
    assert r.status_code == 201 and r.json()["status"] == "confidential", r.text
    name = r.json()["filename"]
    assert name == "Положение об оплате (2).txt"             # скобки в имени сохраняются
    docs = {d["filename"]: d for d in c.get("/documents").json()["documents"]}
    assert docs[name]["confidential"] and docs[name]["status"] == "confidential"
    # тот же файл под другим именем — не обрабатывается повторно
    dup = c.post("/documents/upload", headers=ORIGIN, files={"file": ("copy.txt", body, "text/plain")})
    assert dup.status_code == 200 and dup.json()["duplicate"]
    # переразбор конфиденциального — запрещён; «..» — не документ
    assert c.post(f"/documents/{name}/reprocess", headers=ORIGIN).status_code == 400
    assert c.post("/documents/../reprocess", headers=ORIGIN).status_code in (404, 405)
    assert c.post("/documents/..%2F..%2Fetc/reprocess", headers=ORIGIN).status_code == 404
    # снять флаг — документ уходит на обычную обработку
    off = c.post(f"/documents/{name}/confidential", json={"confidential": False}, headers=ORIGIN).json()
    assert off["status"] == "uploaded" and not off["confidential"]
    assert c.delete(f"/documents/{name}", headers=ORIGIN).status_code == 200


def test_registry_atomic_merge(env):
    import docregistry
    docregistry.update("a.pdf", status="uploaded", folders=["ot"])
    docregistry.update("a.pdf", status="indexed", chunks=5)
    e = docregistry.get("a.pdf")
    assert e["status"] == "indexed" and e["folders"] == ["ot"] and e["chunks"] == 5
    assert docregistry.folder_doc_counts().get("ot") == 1
    docregistry.strip_folder("ot")
    assert docregistry.get("a.pdf")["folders"] == []
    assert docregistry.set_clarification("a.pdf", "устарел") and docregistry.get("a.pdf")["clarification"] == "устарел"
    assert docregistry.remove("a.pdf") and docregistry.get("a.pdf") is None


# ---------------------------------------------------------------- вопросы и чат
def test_questions_in_db_and_order(env, owner):
    import questions
    u = {"id": "emp-q", "full_name": "Вопросов Василий"}
    q1 = questions.record(u, "Первый вопрос", None, questions.REASON_NO_ANSWER)
    questions.record(u, "Пожар в цехе!", None, questions.REASON_ESCALATE, "пожар")
    assert [q["question"] for q in questions.list_for_user("emp-q")][0] == "Первый вопрос"
    assert questions.list_all("open")[0]["reason"] == questions.REASON_ESCALATE
    c = _as_owner(env, owner)
    r = c.post(f"/questions/{q1['id']}/resolve", json={"answer": "Ответ"}, headers=ORIGIN)
    assert r.status_code == 200 and r.json()["status"] == "resolved"


def test_ask_limits(env, owner):
    c = _as_owner(env, owner)
    assert c.post("/ask", json={"question": "x" * 5000}, headers=ORIGIN).status_code == 422
    assert c.post("/ask", json={"question": "   "}, headers=ORIGIN).status_code == 400


def test_push_token_removal_is_scoped(env, owner):
    import push
    import db
    push.register_token("someone-else", "tok-foreign", "android")
    c = _as_owner(env, owner)
    c.request("DELETE", "/api/my/push-token", json={"token": "tok-foreign"}, headers=ORIGIN)
    assert db.query("SELECT 1 FROM push_tokens WHERE token = 'tok-foreign'", (), "one") is not None


def test_ensure_all_ignores_test_notifications(env, owner, monkeypatch):
    """Тестовое уведомление (строка без plan_id) не блокирует материализацию плана."""
    import users
    import messaging
    c = _as_owner(env, owner)
    plan = c.post("/plans/template", headers=ORIGIN).json()
    emp = c.post("/users", headers=ORIGIN, json={"full_name": "Планов Павел Павлович",
                                                 "plan_id": plan["plan_id"], "start_date": "2026-10-01"}).json()
    messaging.deliver_now(emp["id"], "Тест", "Проверка")
    made = []
    monkeypatch.setattr(messaging, "materialize_employee", lambda u, force=False: made.append(u["id"]) or 1)
    messaging.ensure_all()
    assert emp["id"] in made
    assert users.get_user(emp["id"])["start_date"] == "2026-10-01"


def test_answers_size_limit(env, owner):
    c = _as_owner(env, owner)
    r = c.post("/api/my/messages/x/answer", json={"answers": {"a": "x" * 40000}}, headers=ORIGIN)
    assert r.status_code == 413


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
