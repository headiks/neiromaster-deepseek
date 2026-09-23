"""Пользователи: повторная загрузка штатки не плодит дубли, логин из ФИО, раздельные контакты,
автостатус адаптации. Без БД: users для staffing — простая подмена в памяти."""
import sys
import types
from datetime import date

import test_stubs
test_stubs.install(psycopg=True, stub_modules=())
import users as real_users  # noqa: E402
import staffing  # noqa: E402


class FakeUsers:
    ROLE_EMPLOYEE = "employee"

    def __init__(self, rows):
        self.rows = list(rows)

    def list_users(self, with_secrets=False):
        return list(self.rows)

    def create_user(self, raw, role="employee", must_change_credentials=False, **kw):
        row = {"id": f"u{len(self.rows)}", **raw}
        self.rows.append(row)
        return row


def test_reimport_skips_existing_people_and_vacancies(monkeypatch):
    fake = FakeUsers([{"full_name": "Иванов Иван Иванович", "username": "ivanov_i_i"},
                      {"full_name": "(вакансия) Сварщик", "position": "Сварщик", "department": "Цех 1"}])
    monkeypatch.setattr(staffing, "users", fake)
    res = staffing.import_records([
        {"full_name": "Иванов Иван Иванович", "position": "Мастер"},
        {"full_name": "Иванова Инна Игоревна", "position": "Мастер"},
        {"full_name": "", "position": "Сварщик", "department": "Цех 1"},
        {"full_name": "", "position": "Сварщик", "department": "Цех 2"},
    ])
    assert [p["full_name"] for p in res["profiles"]] == ["Иванова Инна Игоревна"]
    assert res["profiles"][0]["username"] == "ivanova_i_i" and res["profiles"][0]["id"]
    assert res["vacancies"] == [{"position": "Сварщик", "department": "Цех 2"}]
    assert len(res["skipped"]) == 2
    again = staffing.import_records([{"full_name": "Иванова Инна Игоревна"}])
    assert again["profiles"] == [] and again["skipped"][0]["reason"] == "уже есть в системе"


def test_credentials_xlsx_has_rows():
    import io
    from openpyxl import load_workbook
    body = staffing.credentials_xlsx([{"full_name": "Петров П.", "username": "petrov_p", "temp_password": "Abc12345"}])
    ws = load_workbook(io.BytesIO(body)).active
    assert [c.value for c in ws[2]][:3] == ["Петров П.", "petrov_p", "Abc12345"]


def test_split_contact():
    assert real_users.split_contact("+7 900 123-45-67, ivan@corp.ru") == ("+7 900 123-45-67", "ivan@corp.ru")
    assert real_users.split_contact("ivan@corp.ru") == ("", "ivan@corp.ru")
    assert real_users.split_contact("@ivan_tg") == ("@ivan_tg", "")


def test_adaptation_status_is_automatic():
    plan = {"stages": [{"anchor": "before_start", "duration": {"value": 3, "unit": "days"}},
                       {"anchor": "from_start", "duration": {"value": 1, "unit": "months"}}]}
    today = date(2026, 9, 23)
    st = lambda **u: real_users.adaptation_status(u, plan, today=today)
    assert st() == "planned"                                   # нет даты выхода
    assert st(start_date="2026-10-01") == "planned"
    assert st(start_date="2026-09-10") == "active"
    assert st(start_date="2026-08-01") == "done"               # 30 дней прошло
    assert st(start_date="2026-09-10", status="paused") == "paused"
