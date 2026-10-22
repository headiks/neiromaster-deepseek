"""
analytics.py — контроль мини-тестов адаптации и аналитика вопросов сотрудников.

Тесты: по каждому доставленному мини-тесту плана — сдан (верных ≥ PASS), не сдан, не пройден
(нет ответов дольше OVERDUE_DAYS) или ждёт ответа. Пересдача — администратор очищает ответы.

Вопросы: каждый вопрос ассистенту и кнопкам «Написать…» пишется в ask_log (текст зашифрован).
Категория — организационный или рабочий вопрос: сначала по ключевым словам, неясные — пачкой
через модель (обезличивание — в deepseek.chat), при недоступности модели остаются «не определено».
"""
import json
import re
from datetime import datetime, timedelta, timezone

import db
import users

PASS = 0.7
OVERDUE_DAYS = 2
CATEGORIES = {"org": "Организационные", "work": "Рабочие"}
OUTCOMES = {"answered": "Ответил ассистент", "escalated": "Передан специалисту",
            "no_answer": "Нет ответа", "to_person": "Написал наставнику / руководителю / в отдел кадров"}

_ORG = re.compile(
    r"зарплат|аванс|оклад|преми|выплат|отпуск|больнич|график|смен[аыуе]\b|выходн|пропуск|проходн|"
    r"столов|обед|питани|спецодежд|\bсиз\b|документ|справк|договор|трудов\w* книж|оформ|кадр|"
    r"медосмотр|медкомисс|страхов|парковк|автобус|транспорт|доехать|добрат|раздевалк|шкафчик|"
    r"общежит|льгот|компенсац|отгул|опоздан|увольн|испытательн|где наход|кому обрати|телефон")
_WORK = re.compile(
    r"станок|станк|оборуд|инструмент|технолог|операци|чертеж|чертёж|детал|наряд|норм[аыу]\b|выработ|"
    r"брак|качеств|сварк|сварн|ремонт|маршрут|рейс|погрузк|разгрузк|материал|сырь|наладк|настрой|"
    r"режим работ|охран\w* труда|техник\w* безопасн|инструктаж|допуск|аттестац|экзамен|обязанност|"
    r"как правильно|порядок работ|процесс|смазк|электр|кран\b|погрузчик|пульт|программ")


# ---------- Тесты ----------
def quiz_status(right: int, answered: int, total: int, delivered_at, now=None) -> str:
    """'passed' | 'failed' | 'overdue' (не пройден) | 'waiting'."""
    if answered and total:
        return "passed" if right / total >= PASS else "failed"
    now = now or datetime.now(timezone.utc)
    if delivered_at and now - delivered_at > timedelta(days=OVERDUE_DAYS):
        return "overdue"
    return "waiting"


def tests(employee_ids: list) -> list:
    """Мини-тесты плана сотрудников: [{id, employee_id, title, delivered_at, right, total, status}]."""
    import stats
    ids = [i for i in set(employee_ids or []) if i]
    if not ids:
        return []
    rows = db.query("SELECT id, employee_id, title, payload, answers, delivered_at FROM scheduled_messages "
                    "WHERE kind = 'quiz' AND plan_id IS NOT NULL AND status IN ('delivered', 'read') "
                    "AND employee_id = ANY(%s) ORDER BY delivered_at", (ids,)) or []
    out = []
    for r in rows:
        right, answered, total = stats.quiz_score(r["payload"], r["answers"])
        out.append({"id": r["id"], "employee_id": r["employee_id"],
                    "title": " · ".join(p for p in str(r["title"] or "").split(" — ") if p),
                    "delivered_at": str(r["delivered_at"] or ""), "right": right, "answered": answered,
                    "total": total, "status": quiz_status(right, answered, total, r["delivered_at"])})
    return out


def retake(row_id: str) -> str | None:
    """Пересдача: ответы теста очищаются. -> employee_id или None (не тест / не найден)."""
    row = db.query("UPDATE scheduled_messages SET answers = NULL, updated_at = now() "
                   "WHERE id = %s AND kind = 'quiz' RETURNING employee_id", (row_id,), "one")
    return row["employee_id"] if row else None


# ---------- Вопросы ----------
def log_question(user: dict, question: str, outcome: str):
    """Сбой журнала не мешает ответу сотруднику."""
    try:
        db.execute("INSERT INTO ask_log (user_id, question, outcome, category) VALUES (%s, %s, %s, %s)",
                   (user.get("id"), users._encrypt_field((question or "").strip()[:2000]),
                    outcome if outcome in OUTCOMES else "answered", classify(question)))
    except Exception as e:
        print(f"[analytics] вопрос не записан: {e}")


def classify(text: str) -> str | None:
    """По ключевым словам: 'org' | 'work' | None (неясно — решит модель)."""
    low = (text or "").lower().replace("ё", "е")
    org, work = len(_ORG.findall(low)), len(_WORK.findall(low))
    return "org" if org > work else "work" if work > org else None


_CLASSIFY_SYSTEM = """Раздели вопросы новых сотрудников завода на две группы:
org — организационные: оформление, документы, зарплата, график, пропуск, питание, транспорт,
спецодежда, льготы, к кому обратиться;
work — рабочие: как выполнять работу, оборудование, технология, нормы, качество, охрана труда
на рабочем месте, обучение профессии.
Верни JSON: {"1": "org", "2": "work", ...} — по номеру каждого вопроса."""


def classify_pending(limit: int = 100) -> int:
    """Вопросы без категории -> модель, пачкой. -> сколько размечено."""
    rows = db.query("SELECT id, question FROM ask_log WHERE category IS NULL ORDER BY id DESC LIMIT %s",
                    (limit,)) or []
    if not rows:
        return 0
    texts = [users._decrypt_field(r["question"]) or "" for r in rows]
    try:
        import deepseek
        raw = deepseek.chat(_CLASSIFY_SYSTEM, "\n".join(f"{i + 1}. {t[:300]}" for i, t in enumerate(texts)),
                            json_mode=True, max_tokens=20 * len(rows) + 50, retries=1)
        picks = json.loads(raw)
    except Exception as e:
        print(f"[analytics] вопросы не размечены моделью: {e}")
        return 0
    n = 0
    for i, r in enumerate(rows):
        cat = picks.get(str(i + 1))
        if cat in CATEGORIES:
            db.execute("UPDATE ask_log SET category = %s WHERE id = %s", (cat, r["id"]))
            n += 1
    return n


def questions(employee_ids: list, days: int = 30) -> dict:
    """Сводка и список вопросов сотрудников за days дней."""
    ids = [i for i in set(employee_ids or []) if i]
    if not ids:
        return {"total": 0, "by_category": {}, "by_outcome": {}, "by_employee": {}, "by_week": {}, "items": []}
    classify_pending()
    rows = db.query("SELECT id, ts, user_id, question, outcome, category FROM ask_log "
                    "WHERE user_id = ANY(%s) AND ts > now() - make_interval(days => %s) ORDER BY ts DESC",
                    (ids, max(1, min(int(days), 365)))) or []
    by_cat, by_out, by_emp, by_week, items = {}, {}, {}, {}, []
    for r in rows:
        cat = r["category"] or "unknown"
        by_cat[cat] = by_cat.get(cat, 0) + 1
        week = str(r["ts"].date() - timedelta(days=r["ts"].weekday()))      # понедельник недели
        w = by_week.setdefault(week, {"org": 0, "work": 0, "unknown": 0})
        w[cat] = w.get(cat, 0) + 1
        by_out[r["outcome"]] = by_out.get(r["outcome"], 0) + 1
        e = by_emp.setdefault(r["user_id"], {"total": 0, "org": 0, "work": 0, "unknown": 0, "unanswered": 0})
        e["total"] += 1
        e[cat] = e.get(cat, 0) + 1
        e["unanswered"] += r["outcome"] in ("escalated", "no_answer")
        if len(items) < 300:
            items.append({"id": r["id"], "ts": str(r["ts"]), "user_id": r["user_id"], "category": cat,
                          "outcome": r["outcome"], "question": users._decrypt_field(r["question"]) or ""})
    return {"total": len(rows), "by_category": by_cat, "by_outcome": by_out, "by_employee": by_emp,
            "by_week": dict(sorted(by_week.items())), "items": items}


# ---------- Сводка, зона риска, отчёт ----------
LOW_ENGAGEMENT = 40


def risks(row: dict) -> list:
    """Почему новичок в зоне риска (пусто — не в зоне): решение наставника, тесты, вовлечённость."""
    out = []
    if (row.get("review") or {}).get("decision") == "risk":
        out.append("Наставник видит риск")
    t = row.get("tests") or {}
    if t.get("failed"):
        out.append(f"Не сдано тестов: {t['failed']}")
    if t.get("overdue"):
        out.append(f"Не пройдено вовремя: {t['overdue']}")
    if row.get("engagement") is not None and row["engagement"] < LOW_ENGAGEMENT:
        out.append("Низкая вовлечённость")
    return out


def tests_by(employee_ids: list) -> tuple:
    """Тесты одной выборкой -> ({employee_id: счётчики статусов}, [по тестам: title + счётчики])."""
    zero = ("given", "passed", "failed", "overdue", "waiting")
    by_emp, by_title = {}, {}
    for t in tests(employee_ids):
        for c in (by_emp.setdefault(t["employee_id"], dict.fromkeys(zero, 0)),
                  by_title.setdefault(t["title"], {"title": t["title"], **dict.fromkeys(zero, 0)})):
            c["given"] += 1
            c[t["status"]] += 1
    return by_emp, list(by_title.values())


def _pct(v):
    return round(v * 100) if v is not None else None


REPORT_COLUMNS = (
    ("ФИО", 32, lambda r: r.get("full_name")),
    ("Должность", 26, lambda r: r.get("position")),
    ("Подразделение", 26, lambda r: r.get("department")),
    ("Наставник", 28, lambda r: r.get("mentor")),
    ("Дата выхода", 12, lambda r: r.get("start_date")),
    ("Этап", 30, lambda r: r.get("stage")),
    ("Прогресс, %", 11, lambda r: _pct(r.get("progress"))),
    ("Тестов выдано", 9, lambda r: r["tests"]["given"]),
    ("Сдано", 8, lambda r: r["tests"]["passed"]),
    ("Не сдано", 8, lambda r: r["tests"]["failed"]),
    ("Не пройдено вовремя", 11, lambda r: r["tests"]["overdue"]),
    ("Верных ответов, %", 11, lambda r: r["tests"].get("percent")),
    ("Вовлечённость, 0–100", 13, lambda r: r.get("engagement")),
    ("Читает", 8, lambda r: (r.get("reading") or {}).get("read")),
    ("Бегло", 8, lambda r: (r.get("reading") or {}).get("skim")),
    ("Пролистывает", 11, lambda r: (r.get("reading") or {}).get("scroll")),
    ("Последний оценённый этап", 30, lambda r: (r.get("review") or {}).get("stage_title")),
    ("Средний балл наставника", 11, lambda r: (r.get("review") or {}).get("avg")),
    ("Решение наставника", 24, lambda r: (r.get("review") or {}).get("decision_label")),
    ("Ждут оценки", 9, lambda r: r.get("pending_reviews")),
    ("Вопросов", 9, lambda r: r["questions"]["total"]),
    ("Организационных", 11, lambda r: r["questions"].get("org", 0)),
    ("Рабочих", 9, lambda r: r["questions"].get("work", 0)),
    ("Без ответа", 9, lambda r: r["questions"].get("unanswered", 0)),
    ("Зона риска", 40, lambda r: "; ".join(r.get("risks") or [])),
)


def report_xlsx(rows: list, scope: str, company: str = "") -> bytes:
    """Excel-отчёт: шапка со сводкой по выборке и строка на новичка (REPORT_COLUMNS)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from staffing import safe_cell
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    avg = lambda vals: round(sum(vals) / len(vals)) if vals else "—"
    given = sum(r["tests"]["given"] for r in rows)
    lines = [
        "Отчёт по адаптации новичков",
        f"Компания: {company}" if company else None,
        f"Выборка: {scope}",
        f"Сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        f"Новичков: {len(rows)} · средний прогресс: {avg([_pct(r['progress']) for r in rows if r.get('progress') is not None])}% · "
        f"средняя вовлечённость: {avg([r['engagement'] for r in rows if r.get('engagement') is not None])} · "
        f"тестов сдано: {sum(r['tests']['passed'] for r in rows)} из {given} · "
        f"в зоне риска: {sum(bool(r.get('risks')) for r in rows)}",
    ]
    for line in filter(None, lines):
        ws.append([safe_cell(line)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append([c[0] for c in REPORT_COLUMNS])
    head = ws.max_row
    for cell in ws[head]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([v if isinstance(v, (int, float)) else safe_cell(v) for v in (c[2](r) for c in REPORT_COLUMNS)])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"         # строка, а не формула (ФИО из загруженной штатки)
    for i, c in enumerate(REPORT_COLUMNS):
        ws.column_dimensions[ws.cell(head, i + 1).column_letter].width = c[1]
    ws.freeze_panes = ws.cell(head + 1, 2)
    ws.auto_filter.ref = f"A{head}:{ws.cell(max(head, ws.max_row), len(REPORT_COLUMNS)).coordinate}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    now = datetime(2026, 1, 10, tzinfo=timezone.utc)
    assert quiz_status(3, 3, 3, now, now) == "passed" and quiz_status(2, 3, 3, now, now) == "failed"
    assert quiz_status(7, 10, 10, now, now) == "passed"                  # ровно порог — сдан
    assert quiz_status(0, 0, 3, now - timedelta(days=3), now) == "overdue"
    assert quiz_status(0, 0, 3, now - timedelta(hours=5), now) == "waiting"
    assert classify("Когда выдают аванс и зарплату?") == "org"
    assert classify("Как правильно настроить станок перед сменой работы?") == "work"
    assert classify("Где получить спецодежду?") == "org"
    assert classify("Что делать при нарушении техники безопасности на участке?") == "work"
    assert classify("Привет") is None
    calm = {"tests": {"failed": 0, "overdue": 0}, "engagement": 80, "review": None}
    assert risks(calm) == []
    assert risks({**calm, "engagement": 20, "review": {"decision": "risk"}}) == ["Наставник видит риск", "Низкая вовлечённость"]
    assert risks({**calm, "tests": {"failed": 1, "overdue": 2}}) == ["Не сдано тестов: 1", "Не пройдено вовремя: 2"]
    try:
        import io
        from openpyxl import load_workbook
    except ImportError:
        print("SKIP report_xlsx: нет openpyxl")
    else:
        row = {**calm, "full_name": "=HYPERLINK(1)", "progress": 0.5, "risks": [],
               "tests": {"given": 2, "passed": 1, "failed": 0, "overdue": 0, "percent": 75},
               "questions": {"total": 3, "org": 1, "work": 2, "unanswered": 0}}
        ws = load_workbook(io.BytesIO(report_xlsx([row], "вся компания"))).active
        cells = [c.value for line in ws.iter_rows() for c in line]
        assert "'=HYPERLINK(1)" in cells and 50 in cells and 75 in cells     # формула не исполняется, числа — числа
    print("OK: analytics")
